"""Conversão Excel → PDF via ReportLab (fallback quando engine != excel_com)."""

from __future__ import annotations

import io

from openpyxl import load_workbook
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas

from .report_excel_export import format_cell_value_for_display


def convert_workbook_bytes_to_pdf(xlsx_bytes: bytes) -> bytes:
    workbook = load_workbook(io.BytesIO(xlsx_bytes))
    pdf_buffer = io.BytesIO()
    pdf = canvas.Canvas(pdf_buffer, pagesize=landscape(A4))

    page_width, page_height = landscape(A4)
    margin = 20 * mm
    cell_height = 6 * mm
    cell_width = 25 * mm

    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]

        pdf.setFont("Helvetica-Bold", 14)
        pdf.drawString(margin, page_height - margin, f"Planilha: {sheet_name}")

        y_position = page_height - margin - 20 * mm
        pdf.setFont("Helvetica", 9)

        max_row = sheet.max_row
        max_col = sheet.max_column
        rows_per_page = int((page_height - 2 * margin - 20 * mm) / cell_height)
        cols_per_page = int((page_width - 2 * margin) / cell_width)

        for row_offset in range(0, max_row, rows_per_page):
            if row_offset > 0:
                pdf.showPage()
                pdf.setFont("Helvetica-Bold", 14)
                pdf.drawString(
                    margin,
                    page_height - margin,
                    f"Planilha: {sheet_name} (continuação)",
                )
                y_position = page_height - margin - 20 * mm
                pdf.setFont("Helvetica", 9)

            for row_idx in range(1, min(rows_per_page + 1, max_row - row_offset + 1)):
                actual_row = row_offset + row_idx
                x_position = margin

                for col_idx in range(1, min(cols_per_page + 1, max_col + 1)):
                    cell = sheet.cell(row=actual_row, column=col_idx)
                    cell_value = format_cell_value_for_display(cell)

                    pdf.rect(x_position, y_position, cell_width, cell_height)

                    if len(cell_value) > 15:
                        cell_value = cell_value[:12] + "..."

                    pdf.drawString(x_position + 2, y_position + 2, cell_value)
                    x_position += cell_width

                y_position -= cell_height

        if sheet_name != workbook.sheetnames[-1]:
            pdf.showPage()

    pdf.save()
    return pdf_buffer.getvalue()
