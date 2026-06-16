"""
Teste de geração do template Excel
"""
import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

try:
    print("Criando workbook...")
    wb = Workbook()

    print("Configurando planilha Dados...")
    ws_dados = wb.active
    ws_dados.title = "Dados"

    # Cabeçalhos
    headers = [
        "Página", "Ordem Página", "Campo", "Tipo", "Label",
        "Ordem Campo", "Obrigatório", "Config Adicional", "Excel Mapping"
    ]
    ws_dados.append(headers)

    # Estilizar cabeçalhos
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_alignment = Alignment(horizontal="center", vertical="center")

    for col_num, header in enumerate(headers, 1):
        cell = ws_dados.cell(1, col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    print("Adicionando exemplos...")
    examples = [
        ["Dados Gerais", 1, "nome_completo", "textbox", "Nome Completo", 1, "SIM",
         '{"max_characters": 100}', '[{"planilha":"Dados","celula":"A5"}]'],
    ]

    for row in examples:
        ws_dados.append(row)

    print("Criando planilha Instruções...")
    ws_instrucoes = wb.create_sheet("Instruções")
    ws_instrucoes.cell(1, 1, "TESTE")

    print("Salvando em memória...")
    excel_buffer = io.BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)

    print(f"✅ Excel gerado com sucesso! Tamanho: {len(excel_buffer.getvalue())} bytes")

except Exception as e:
    print(f"❌ ERRO: {e}")
    import traceback
    traceback.print_exc()
